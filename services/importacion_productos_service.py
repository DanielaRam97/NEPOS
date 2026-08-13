from dataclasses import dataclass, field

from openpyxl import load_workbook
from database.conexion import nueva_sesion
from database.modelos import Categoria, Producto, Proveedor, Usuario
from services.producto_service import ErrorProducto, ROLES_GESTION_PRODUCTOS
from services.iva_service import iva_categoria_formateado


MODOS_IMPORTACION = {
    "SOLO_NUEVOS": "Solo productos nuevos",
    "ACTUALIZAR": "Actualizar datos sin modificar stock",
    "ACTUALIZAR_SUMAR_STOCK": "Actualizar datos y sumar stock",
}

ENCABEZADOS = {
    "CODIGO": ("CODIGO", "CÓDIGO", "COD DE BARRAS", "CODIGO DE BARRAS"),
    "PRODUCTO": ("PRODUCTO", "DESCRIPCION", "DESCRIPCIÓN"),
    "CATEGORIA": ("CATEGORIA", "CATEGORÍA"),
    "PROVEEDOR": ("PROVEEDOR",),
    "COSTO": ("COSTO",),
    "INCREMENTO": ("INCREMENTO", "MARGEN", "MARGEN %"),
    "STOCK": ("STOCK", "STOCK ACTUAL", "STOCK INICIAL"),
    "IVA": ("IVA",),
    "PESABLE": ("PESABLE",),
    "PLU": ("PLU",),
}


@dataclass
class ResultadoAnalisis:
    filas_validas: list[dict] = field(default_factory=list)
    errores: list[str] = field(default_factory=list)
    total_filas: int = 0
    nuevos: int = 0
    existentes: int = 0
    proveedores_nuevos: set[str] = field(default_factory=set)
    categorias_nuevas: set[str] = field(default_factory=set)


def _texto(valor):
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def _normalizar_encabezado(valor):
    return " ".join(_texto(valor).upper().split())


def _indice_encabezados(fila):
    disponibles = {_normalizar_encabezado(valor): indice for indice, valor in enumerate(fila)}
    indices = {}
    for campo, aliases in ENCABEZADOS.items():
        indices[campo] = next(
            (disponibles[alias] for alias in aliases if alias in disponibles),
            None,
        )
    faltantes = [campo for campo in ("CODIGO", "PRODUCTO", "CATEGORIA", "COSTO") if indices[campo] is None]
    if faltantes:
        raise ErrorProducto(
            "Faltan columnas obligatorias: " + ", ".join(faltantes)
        )
    return indices


def _valor(fila, indices, campo):
    indice = indices.get(campo)
    return fila[indice] if indice is not None and indice < len(fila) else None


def _numero(valor, nombre, numero_fila, por_defecto=0.0):
    if valor in (None, ""):
        return por_defecto
    try:
        numero = float(str(valor).replace("$", "").replace("%", "").replace(",", ".").strip())
    except (TypeError, ValueError):
        raise ErrorProducto(f"Fila {numero_fila}: {nombre} debe ser numérico.")
    if numero < 0:
        raise ErrorProducto(f"Fila {numero_fila}: {nombre} no puede ser negativo.")
    return numero


def _booleano(valor):
    return _texto(valor).upper() in ("SI", "SÍ", "TRUE", "VERDADERO", "1", "X")


def _validar_permiso(db, usuario_id):
    usuario = db.get(Usuario, usuario_id)
    if not usuario or not usuario.activo:
        raise ErrorProducto("Usuario inválido o inactivo.")
    if usuario.rol not in ROLES_GESTION_PRODUCTOS:
        raise ErrorProducto("No tenés permiso para importar productos.")


def analizar_excel(ruta):
    try:
        libro = load_workbook(ruta, read_only=True, data_only=True)
    except Exception as error:
        raise ErrorProducto(f"No se pudo abrir el Excel: {error}")

    hoja = libro.active
    filas = hoja.iter_rows(values_only=True)
    encabezado = next(filas, None)
    if not encabezado:
        libro.close()
        raise ErrorProducto("El archivo Excel está vacío.")

    indices = _indice_encabezados(encabezado)
    resultado = ResultadoAnalisis()
    codigos_archivo = set()
    plus_archivo = set()

    db = nueva_sesion()
    try:
        categorias = {
            categoria.nombre.strip().upper(): categoria
            for categoria in db.query(Categoria).all()
        }
        proveedores = {
            proveedor.nombre.strip().upper(): proveedor
            for proveedor in db.query(Proveedor).all()
        }
        codigos_existentes = {
            codigo for (codigo,) in db.query(Producto.codigo).all()
        }
        plus_existentes = {
            plu for (plu,) in db.query(Producto.plu).filter(Producto.plu.isnot(None)).all()
        }

        for numero_fila, fila in enumerate(filas, start=2):
            if not fila or all(valor in (None, "") for valor in fila):
                continue
            resultado.total_filas += 1

            try:
                codigo = _texto(_valor(fila, indices, "CODIGO"))
                descripcion = _texto(_valor(fila, indices, "PRODUCTO"))
                categoria_nombre = _texto(_valor(fila, indices, "CATEGORIA")).upper()
                proveedor_nombre = _texto(_valor(fila, indices, "PROVEEDOR")).upper()
                plu = _texto(_valor(fila, indices, "PLU")) or None

                if not codigo:
                    raise ErrorProducto(f"Fila {numero_fila}: el código es obligatorio.")
                if not descripcion:
                    raise ErrorProducto(f"Fila {numero_fila}: el producto es obligatorio.")
                if not categoria_nombre:
                    raise ErrorProducto(f"Fila {numero_fila}: la categoría es obligatoria.")
                if codigo in codigos_archivo:
                    raise ErrorProducto(f"Fila {numero_fila}: código repetido dentro del Excel ({codigo}).")
                if plu and plu in plus_archivo:
                    raise ErrorProducto(f"Fila {numero_fila}: PLU repetido dentro del Excel ({plu}).")

                categoria = categorias.get(categoria_nombre)
                if not categoria:
                    resultado.categorias_nuevas.add(categoria_nombre)
                proveedor = proveedores.get(proveedor_nombre) if proveedor_nombre else None
                if proveedor_nombre and not proveedor:
                    resultado.proveedores_nuevos.add(proveedor_nombre)

                costo = _numero(_valor(fila, indices, "COSTO"), "el costo", numero_fila)
                incremento = _numero(
                    _valor(fila, indices, "INCREMENTO"), "el incremento", numero_fila
                )
                if incremento > 1:
                    incremento /= 100
                stock = _numero(_valor(fila, indices, "STOCK"), "el stock", numero_fila)
                pesable = _booleano(_valor(fila, indices, "PESABLE"))

                if plu and plu in plus_existentes and codigo not in codigos_existentes:
                    raise ErrorProducto(
                        f"Fila {numero_fila}: el PLU '{plu}' ya pertenece a otro producto."
                    )

                codigos_archivo.add(codigo)
                if plu:
                    plus_archivo.add(plu)

                es_existente = codigo in codigos_existentes
                resultado.existentes += 1 if es_existente else 0
                resultado.nuevos += 0 if es_existente else 1
                resultado.filas_validas.append({
                    "fila": numero_fila,
                    "codigo": codigo,
                    "descripcion": descripcion,
                    "categoria_id": categoria.id if categoria else None,
                    "categoria_nombre": categoria_nombre,
                    "proveedor_id": proveedor.id if proveedor else None,
                    "proveedor_nombre": proveedor_nombre or None,
                    "costo": costo,
                    "incremento": incremento,
                    "stock": stock,
                    "iva": _texto(_valor(fila, indices, "IVA")) or (
                        categoria.iva if categoria else "21%"
                    ),
                    "pesable": pesable,
                    "plu": plu,
                    "existente": es_existente,
                })
            except ErrorProducto as error:
                resultado.errores.append(str(error))
    finally:
        db.close()
        libro.close()

    return resultado


def importar_excel(ruta, modo, usuario_id):
    if modo not in MODOS_IMPORTACION:
        raise ErrorProducto("Modo de importación inválido.")

    analisis = analizar_excel(ruta)
    if analisis.errores:
        raise ErrorProducto("Corregí los errores del Excel antes de importar.")

    db = nueva_sesion()
    creados = 0
    actualizados = 0
    omitidos = 0
    try:
        _validar_permiso(db, usuario_id)
        proveedores = {
            proveedor.nombre.strip().upper(): proveedor
            for proveedor in db.query(Proveedor).all()
        }
        categorias = {
            categoria.nombre.strip().upper(): categoria
            for categoria in db.query(Categoria).all()
        }
        for datos in analisis.filas_validas:
            categoria = categorias.get(datos["categoria_nombre"])
            if not categoria:
                categoria = Categoria(
                    nombre=datos["categoria_nombre"],
                    iva=iva_categoria_formateado(datos["categoria_nombre"]),
                )
                db.add(categoria)
                db.flush()
                categorias[datos["categoria_nombre"]] = categoria

            proveedor = None
            if datos["proveedor_nombre"]:
                proveedor = proveedores.get(datos["proveedor_nombre"])
                if not proveedor:
                    proveedor = Proveedor(nombre=datos["proveedor_nombre"])
                    db.add(proveedor)
                    db.flush()
                    proveedores[datos["proveedor_nombre"]] = proveedor

            producto = (
                db.query(Producto)
                .filter(Producto.codigo == datos["codigo"])
                .with_for_update()
                .first()
            )

            if producto and modo == "SOLO_NUEVOS":
                omitidos += 1
                continue

            if producto:
                producto.descripcion = datos["descripcion"]
                producto.categoria_id = categoria.id
                producto.proveedor_id = proveedor.id if proveedor else None
                producto.costo = datos["costo"]
                producto.incremento = datos["incremento"]
                producto.precio = round(datos["costo"] * (1 + datos["incremento"]), 2)
                producto.iva = iva_categoria_formateado(categoria.nombre)
                producto.pesable = datos["pesable"]
                producto.plu = datos["plu"]
                if modo == "ACTUALIZAR_SUMAR_STOCK":
                    producto.stock = (producto.stock or 0) + datos["stock"]
                actualizados += 1
            else:
                db.add(Producto(
                    codigo=datos["codigo"],
                    plu=datos["plu"],
                    descripcion=datos["descripcion"],
                    categoria_id=categoria.id,
                    proveedor_id=proveedor.id if proveedor else None,
                    costo=datos["costo"],
                    incremento=datos["incremento"],
                    precio=round(datos["costo"] * (1 + datos["incremento"]), 2),
                    stock=datos["stock"],
                    iva=iva_categoria_formateado(categoria.nombre),
                    pesable=datos["pesable"],
                    activo=1,
                ))
                creados += 1

        db.commit()
        return {"creados": creados, "actualizados": actualizados, "omitidos": omitidos}
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()
