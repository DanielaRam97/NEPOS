from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

ENCABEZADOS = ["N°", "Fecha", "Cajero", "Turno", "Forma de pago", "Total", "Estado"]


def _fila_de_venta(venta):
    cajero = venta.usuario.nombre if venta.usuario else "-"
    turno = venta.turno_rel.turno if venta.turno_rel else "-"
    estado = "Anulada" if venta.anulada else "Activa"
    return [
        venta.numero or venta.id,
        venta.fecha.strftime("%d/%m/%Y %H:%M"),
        cajero,
        turno,
        venta.forma_pago or "-",
        float(venta.total),
        estado,
    ]


def exportar_ventas_excel(ventas, ruta_archivo):
    wb = Workbook()
    ws = wb.active
    ws.title = "Historial de ventas"

    ws.append(ENCABEZADOS)
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1E293B")

    for venta in ventas:
        ws.append(_fila_de_venta(venta))

    for columna in ws.columns:
        largo_max = max((len(str(c.value)) for c in columna if c.value is not None), default=10)
        ws.column_dimensions[columna[0].column_letter].width = largo_max + 2

    wb.save(ruta_archivo)


def exportar_ventas_pdf(ventas, ruta_archivo, titulo="Historial de ventas"):
    doc = SimpleDocTemplate(ruta_archivo, pagesize=landscape(A4))
    estilos = getSampleStyleSheet()
    elementos = [Paragraph(titulo, estilos["Title"]), Spacer(1, 0.5 * cm)]

    filas = [ENCABEZADOS]
    for venta in ventas:
        numero, fecha, cajero, turno, forma_pago, total, estado = _fila_de_venta(venta)
        filas.append([str(numero), fecha, cajero, turno, forma_pago, f"${total:.2f}", estado])

    tabla = Table(filas, repeatRows=1)
    tabla.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e293b")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f5f9")]),
    ]))
    elementos.append(tabla)
    doc.build(elementos)

def exportar_cierre_turno_excel(resumen, ruta_archivo):
    wb = Workbook()
    ws = wb.active
    ws.title = "Cierre de turno"

    filas = [
        ("Turno", resumen["turno"]),
        ("Cajero", resumen["cajero"]),
        ("Transacciones", resumen["transacciones"]),
        ("Unidades vendidas", resumen["unidades"]),
        ("Pesables (kg)", resumen["pesables"]),
        ("Venta bruta", resumen["venta_bruta"]),
        ("Recargos", resumen["recargos"]),
        ("Neto gravado 10,5%", resumen.get("neto_10_5", 0)),
        ("IVA 10,5% incluido", resumen["iva_10_5"]),
        ("Neto gravado 21%", resumen.get("neto_21", 0)),
        ("IVA 21% incluido", resumen["iva_21"]),
        ("Efectivo", resumen["efectivo"]),
        ("QR", resumen["qr"]),
        ("Débito", resumen["debito"]),
        ("Crédito", resumen["credito"]),
        ("Fondo inicial", resumen["fondo_inicial"]),
        ("Efectivo esperado en caja", resumen["efectivo_esperado_en_caja"]),
        ("Efectivo contado", resumen.get("efectivo_declarado")),
        ("Diferencia de caja", resumen.get("diferencia_caja")),
        ("Observación", resumen.get("observacion_cierre", "")),
    ]
    for fila in filas:
        ws.append(fila)
    for celda in ws["A"]:
        celda.font = Font(bold=True)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 20

    wb.save(ruta_archivo)


def exportar_cierre_turno_pdf(resumen, ruta_archivo):
    doc = SimpleDocTemplate(ruta_archivo, pagesize=A4)
    estilos = getSampleStyleSheet()
    elementos = [
        Paragraph(f"Cierre de turno — {resumen['turno']}", estilos["Title"]),
        Spacer(1, 0.5 * cm),
    ]

    filas = [
        ["Cajero", str(resumen["cajero"])],
        ["Transacciones", str(resumen["transacciones"])],
        ["Unidades vendidas", str(resumen["unidades"])],
        ["Pesables (kg)", str(resumen["pesables"])],
        ["Venta bruta", f"${resumen['venta_bruta']:.2f}"],
        ["Recargos", f"${resumen['recargos']:.2f}"],
        ["Neto gravado 10,5%", f"${resumen.get('neto_10_5', 0):.2f}"],
        ["IVA 10,5% incluido", f"${resumen['iva_10_5']:.2f}"],
        ["Neto gravado 21%", f"${resumen.get('neto_21', 0):.2f}"],
        ["IVA 21% incluido", f"${resumen['iva_21']:.2f}"],
        ["Efectivo", f"${resumen['efectivo']:.2f}"],
        ["QR", f"${resumen['qr']:.2f}"],
        ["Débito", f"${resumen['debito']:.2f}"],
        ["Crédito", f"${resumen['credito']:.2f}"],
        ["Fondo inicial", f"${resumen['fondo_inicial']:.2f}"],
        ["Efectivo esperado en caja", f"${resumen['efectivo_esperado_en_caja']:.2f}"],
        [
            "Efectivo contado",
            (
                f"${resumen['efectivo_declarado']:.2f}"
                if resumen.get("efectivo_declarado") is not None
                else "—"
            ),
        ],
        [
            "Diferencia de caja",
            (
                f"${resumen['diferencia_caja']:.2f}"
                if resumen.get("diferencia_caja") is not None
                else "—"
            ),
        ],
        ["Observación", resumen.get("observacion_cierre", "") or "—"],
    ]
    tabla = Table(filas)
    tabla.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#f1f5f9")),
    ]))
    elementos.append(tabla)
    doc.build(elementos)
