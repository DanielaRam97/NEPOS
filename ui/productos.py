from datetime import datetime
from utils.rutas import ruta_icono

from PySide6.QtCore import QSize, Qt
from PySide6.QtGui import QPixmap
from PySide6.QtWidgets import (
    QCheckBox, QComboBox, QFrame, QGridLayout, QHBoxLayout, QHeaderView, QLabel,
    QFileDialog, QInputDialog, QLineEdit, QMessageBox, QPushButton, QTableWidget, QTableWidgetItem,
    QDialog, QDialogButtonBox, QFormLayout, QStyle, QVBoxLayout, QWidget,
)

from services.producto_service import (
    ErrorProducto,
    actualizar_producto,
    cambiar_estado_producto,
    crear_categoria,
    crear_producto,
    listar_categorias,
    listar_productos,
    listar_proveedores,
)
from services.grupo_precio_service import (
    ErrorGrupoPrecio,
    crear_grupo_precio,
    listar_grupos_precio,
)
from utils.formato import formatear_cantidad, formatear_moneda
from utils.validacion import convertir_decimal_finito
from services.importacion_productos_service_v3 import (
    MODOS_IMPORTACION, analizar_excel, importar_excel,
)
from services.producto_pendiente_service import (
    ErrorProductoPendiente,
    aprobar_producto,
)
from services.permisos import (
    puede_editar_productos,
    puede_exportar_productos,
    puede_importar_productos,
    puede_ver_costos_productos,
)
from ui.estilos import ESTILO_GLOBAL, configurar_boton_primario


ESTILO_PRODUCTOS = """
QPushButton#nuevo_producto {
    background-color: #12aa52;
    color: #ffffff;
    border: 1px solid #0e9849;
    border-radius: 8px;
    padding: 9px 17px;
    font-size: 14px;
    font-weight: 900;
}
QPushButton#nuevo_producto:hover {
    background-color: #0d9245;
}
QPushButton#nuevo_producto:pressed {
    background-color: #087b3a;
}

QPushButton#importar_productos {
    background-color: #887bd4;
    color: #ffffff;
    border: 1px solid #6e61bd;
    border-radius: 8px;
    padding: 9px 15px;
    font-weight: 800;
}
QPushButton#importar_productos:hover {
    background-color: #7568c5;
}
QPushButton#importar_productos:pressed {
    background-color: #6255ae;
}

QPushButton#exportar_productos {
    background-color: #4da8da;
    color: #ffffff;
    border: 1px solid #3189b9;
    border-radius: 8px;
    padding: 9px 15px;
    font-weight: 800;
}
QPushButton#exportar_productos:hover {
    background-color: #3797cc;
}
QPushButton#exportar_productos:pressed {
    background-color: #287ca9;
}

QPushButton#buscar_productos {
    background-color: #4da8da;
    color: #ffffff;
    border: 1px solid #3189b9;
    border-radius: 7px;
    padding: 8px 14px;
    font-weight: 800;
}
QPushButton#buscar_productos:hover {
    background-color: #3797cc;
}
QPushButton#buscar_productos:pressed {
    background-color: #287ca9;
}

QPushButton#limpiar_productos {
    background-color: #e2e8f0;
    color: #334155;
    border: 1px solid #cbd5e1;
    border-radius: 7px;
    padding: 8px 14px;
    font-weight: 700;
}
QPushButton#limpiar_productos:hover {
    background-color: #d6dde6;
}
QPushButton#limpiar_productos:pressed {
    background-color: #c8d1dc;
}

QPushButton#importar_productos:disabled,
QPushButton#exportar_productos:disabled,
QPushButton#nuevo_producto:disabled {
    background-color: #e5e7eb;
    color: #9ca3af;
    border-color: #d4d8de;
}
"""


ESTILO_BOTON_IMPORTAR = """  # Estilo directo: violeta
QPushButton {
    background-color: #887bd4;
    color: #ffffff;
    border: 1px solid #6e61bd;
    border-radius: 8px;
    padding: 9px 16px;
    font-family: "Segoe UI";
    font-size: 13px;
    font-weight: 800;
}
QPushButton:hover { background-color: #7568c5; }
QPushButton:pressed { background-color: #6255ae; }
QPushButton:disabled {
    background-color: #e5e7eb;
    color: #9ca3af;
    border-color: #d4d8de;
}
"""

ESTILO_BOTON_EXPORTAR = """
QPushButton {
    background-color: #4da8da;
    color: #ffffff;
    border: 1px solid #3189b9;
    border-radius: 8px;
    padding: 9px 16px;
    font-family: "Segoe UI";
    font-size: 13px;
    font-weight: 800;
}
QPushButton:hover { background-color: #3797cc; }
QPushButton:pressed { background-color: #287ca9; }
QPushButton:disabled {
    background-color: #e5e7eb;
    color: #9ca3af;
    border-color: #d4d8de;
}
"""

ESTILO_BOTON_NUEVO = """
QPushButton {
    background-color: #12aa52;
    color: #ffffff;
    border: 1px solid #0e9849;
    border-radius: 8px;
    padding: 9px 17px;
    font-family: "Segoe UI";
    font-size: 14px;
    font-weight: 900;
}
QPushButton:hover { background-color: #0d9245; }
QPushButton:pressed { background-color: #087b3a; }
QPushButton:disabled {
    background-color: #e5e7eb;
    color: #9ca3af;
    border-color: #d4d8de;
}
"""

ESTILO_BOTON_BUSCAR = """
QPushButton {
    background-color: #4da8da;
    color: #ffffff;
    border: 1px solid #3189b9;
    border-radius: 7px;
    padding: 8px 14px;
    font-family: "Segoe UI";
    font-size: 13px;
    font-weight: 800;
}
QPushButton:hover { background-color: #3797cc; }
QPushButton:pressed { background-color: #287ca9; }
"""

ESTILO_BOTON_LIMPIAR = """
QPushButton {
    background-color: #e2e8f0;
    color: #334155;
    border: 1px solid #cbd5e1;
    border-radius: 7px;
    padding: 8px 14px;
    font-family: "Segoe UI";
    font-size: 13px;
    font-weight: 700;
}
QPushButton:hover { background-color: #d6dde6; }
QPushButton:pressed { background-color: #c8d1dc; }
"""


class DialogoNuevaCategoria(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Nueva categoría")
        self.setModal(True)
        self.setMinimumWidth(390)

        layout = QVBoxLayout(self)
        descripcion = QLabel(
            "Ingresá el nombre y elegí la alícuota que se aplicará "
            "a los productos de esta categoría."
        )
        descripcion.setWordWrap(True)
        layout.addWidget(descripcion)

        formulario = QFormLayout()
        self.campo_nombre = QLineEdit()
        self.campo_nombre.setPlaceholderText("Ej.: Carnicería")
        self.combo_iva = QComboBox()
        self.combo_iva.addItem("21%", 21.0)
        self.combo_iva.addItem("10,5%", 10.5)
        formulario.addRow("Nombre:", self.campo_nombre)
        formulario.addRow("IVA:", self.combo_iva)
        layout.addLayout(formulario)

        botones = QDialogButtonBox(
            QDialogButtonBox.StandardButton.Save
            | QDialogButtonBox.StandardButton.Cancel
        )
        botones.button(
            QDialogButtonBox.StandardButton.Save
        ).setText("Crear categoría")
        botones.button(
            QDialogButtonBox.StandardButton.Cancel
        ).setText("Cancelar")
        botones.accepted.connect(self.accept)
        botones.rejected.connect(self.reject)
        layout.addWidget(botones)
        self.campo_nombre.returnPressed.connect(self.accept)
        self.campo_nombre.setFocus()

    def datos(self):
        return self.campo_nombre.text(), self.combo_iva.currentData()


class ItemNumerico(QTableWidgetItem):
    """Ordena importes, porcentajes y cantidades por su valor real."""

    def __init__(self, texto, valor):
        super().__init__(texto)
        self.valor_orden = float(valor or 0)

    def __lt__(self, otro):
        if isinstance(otro, ItemNumerico):
            return self.valor_orden < otro.valor_orden
        return super().__lt__(otro)


class VentanaProductos(QWidget):
    COLUMNAS = [
        "Código", "Producto", "Categoría", "Proveedor", "Costo",
        "Precio venta", "Margen", "Stock", "Grupo", "PLU",
        "Activo", "Revisión",
    ]

    def __init__(self, usuario):
        super().__init__()
        self.usuario = usuario
        self.productos = []
        self.producto_id_seleccionado = None
        self.creando_producto = False
        self._columna_orden = None
        self._orden_actual = Qt.SortOrder.AscendingOrder
        self.puede_editar = puede_editar_productos(usuario)
        self.puede_importar = puede_importar_productos(usuario)
        self.puede_exportar = puede_exportar_productos(usuario)
        self.puede_ver_costos = puede_ver_costos_productos(usuario)

        self.setWindowTitle("NEPOS — Productos")
        self.resize(1250, 780)
        self.setStyleSheet(ESTILO_GLOBAL + ESTILO_PRODUCTOS)

        self._armar_interfaz()
        self._cargar_catalogos()
        self.buscar()

    def _armar_interfaz(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(18, 12, 18, 16)
        layout.setSpacing(10)

        cabecera_frame = QFrame()
        cabecera_frame.setObjectName("tarjeta")
        cabecera = QHBoxLayout(cabecera_frame)
        cabecera.setContentsMargins(18, 10, 18, 10)
        cabecera.setSpacing(12)

        icono = QLabel()
        pixmap = QPixmap(str(ruta_icono("productos.png")))
        if not pixmap.isNull():
            icono.setPixmap(
                pixmap.scaled(
                    34,
                    34,
                    Qt.AspectRatioMode.KeepAspectRatio,
                    Qt.TransformationMode.SmoothTransformation,
                )
            )
        else:
            icono.setText("▰")
            icono.setStyleSheet("font-size:25px; color:#4da8da;")
        icono.setFixedSize(40, 40)
        icono.setAlignment(Qt.AlignmentFlag.AlignCenter)
        cabecera.addWidget(icono)

        textos_cabecera = QVBoxLayout()
        textos_cabecera.setSpacing(2)
        titulo = QLabel("PRODUCTOS")
        titulo.setObjectName("titulo_modulo")
        subtitulo = QLabel(
            "Consulta, precios, grupos y revisión del catálogo."
        )
        subtitulo.setObjectName("subtitulo")
        textos_cabecera.addWidget(titulo)
        textos_cabecera.addWidget(subtitulo)
        cabecera.addLayout(textos_cabecera)
        cabecera.addStretch()

        self.boton_importar = QPushButton("IMPORTAR EXCEL")
        self.boton_importar.setObjectName("importar_productos")
        self.boton_importar.setMinimumSize(170, 44)
        self.boton_importar.setIcon(
            self.style().standardIcon(QStyle.StandardPixmap.SP_ArrowDown)
        )
        self.boton_importar.setIconSize(QSize(24, 24))
        self.boton_importar.setStyleSheet(ESTILO_BOTON_IMPORTAR)
        self.boton_importar.clicked.connect(self.importar_desde_excel)
        self.boton_importar.setVisible(self.puede_importar)
        cabecera.addWidget(self.boton_importar)

        self.boton_exportar = QPushButton("EXPORTAR EXCEL")
        self.boton_exportar.setObjectName("exportar_productos")
        self.boton_exportar.setMinimumSize(175, 44)
        self.boton_exportar.setIcon(
            self.style().standardIcon(QStyle.StandardPixmap.SP_ArrowUp)
        )
        self.boton_exportar.setIconSize(QSize(24, 24))
        self.boton_exportar.setStyleSheet(ESTILO_BOTON_EXPORTAR)
        self.boton_exportar.clicked.connect(self.exportar_a_excel)
        self.boton_exportar.setVisible(self.puede_exportar)
        cabecera.addWidget(self.boton_exportar)

        self.boton_nuevo_producto = QPushButton("NUEVO PRODUCTO")
        self.boton_nuevo_producto.setObjectName("nuevo_producto")
        self.boton_nuevo_producto.setMinimumSize(205, 44)
        self.boton_nuevo_producto.setIcon(
            self.style().standardIcon(
                QStyle.StandardPixmap.SP_FileDialogNewFolder
            )
        )
        self.boton_nuevo_producto.setIconSize(QSize(22, 22))
        self.boton_nuevo_producto.setStyleSheet(ESTILO_BOTON_NUEVO)
        self.boton_nuevo_producto.clicked.connect(self.nuevo_producto)
        self.boton_nuevo_producto.setVisible(self.puede_editar)
        cabecera.addWidget(self.boton_nuevo_producto)
        layout.addWidget(cabecera_frame)

        filtros_frame = QFrame()
        filtros_frame.setObjectName("tarjeta")
        filtros = QHBoxLayout(filtros_frame)
        filtros.setContentsMargins(12, 10, 12, 10)
        filtros.setSpacing(8)
        self.campo_busqueda = QLineEdit()
        self.campo_busqueda.setPlaceholderText("Buscar por código, PLU o nombre")
        self.campo_busqueda.setMinimumHeight(40)
        self.campo_busqueda.returnPressed.connect(self.buscar)

        self.combo_categoria_filtro = QComboBox()
        self.combo_categoria_filtro.setMinimumHeight(40)
        self.combo_estado = QComboBox()
        self.combo_estado.setMinimumHeight(40)
        self.combo_estado.addItems(["TODOS", "ACTIVOS", "INACTIVOS"])
        if not self.puede_editar:
            self.combo_estado.setCurrentText("ACTIVOS")
            self.combo_estado.setEnabled(False)

        boton_buscar = QPushButton("BUSCAR")
        boton_buscar.setObjectName("buscar_productos")
        boton_buscar.setMinimumHeight(40)
        boton_buscar.setIcon(
            self.style().standardIcon(QStyle.StandardPixmap.SP_FileDialogContentsView)
        )
        boton_buscar.setIconSize(QSize(20, 20))
        boton_buscar.setStyleSheet(ESTILO_BOTON_BUSCAR)
        boton_buscar.clicked.connect(self.buscar)
        boton_limpiar = QPushButton("LIMPIAR FILTROS")
        boton_limpiar.setObjectName("limpiar_productos")
        boton_limpiar.setMinimumHeight(40)
        boton_limpiar.setStyleSheet(ESTILO_BOTON_LIMPIAR)
        boton_limpiar.clicked.connect(self.limpiar_filtros)

        filtros.addWidget(self.campo_busqueda, 2)
        filtros.addWidget(self.combo_categoria_filtro)
        filtros.addWidget(self.combo_estado)
        filtros.addWidget(boton_buscar)
        filtros.addWidget(boton_limpiar)
        layout.addWidget(filtros_frame)

        self.etiqueta_resultados = QLabel("")
        layout.addWidget(self.etiqueta_resultados)

        self.tabla = QTableWidget(0, 12)
        self.tabla.setAlternatingRowColors(True)
        self.tabla.setHorizontalHeaderLabels(
            [f"{nombre}  ↕" for nombre in self.COLUMNAS]
        )
        self.tabla.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)
        self.tabla.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        self.tabla.horizontalHeader().setSectionResizeMode(1, QHeaderView.ResizeMode.Stretch)
        encabezado = self.tabla.horizontalHeader()
        encabezado.setSectionsClickable(True)
        encabezado.setSortIndicatorShown(True)
        encabezado.setToolTip(
            "Hacé clic en una columna para ordenar. "
            "Volvé a hacer clic para invertir el orden."
        )
        encabezado.sectionClicked.connect(self._ordenar_por_columna)
        self.tabla.setSortingEnabled(False)
        self.tabla.itemSelectionChanged.connect(self.cargar_seleccion)
        if not self.puede_ver_costos:
            for columna_interna in (3, 4, 6, 8):
                self.tabla.setColumnHidden(columna_interna, True)
        layout.addWidget(self.tabla, 2)

        self.panel_edicion = QFrame()
        self.panel_edicion.setObjectName("tarjeta")
        panel_layout = QVBoxLayout(self.panel_edicion)
        panel_layout.setContentsMargins(16, 12, 16, 12)
        panel_layout.setSpacing(7)
        panel_layout.addWidget(
            QLabel("<b>Edición del producto seleccionado</b>")
        )

        form = QGridLayout()
        form.setHorizontalSpacing(12)
        form.setVerticalSpacing(7)

        self.campo_codigo = QLineEdit()
        self.campo_plu = QLineEdit()
        self.campo_descripcion = QLineEdit()
        self.combo_categoria = QComboBox()
        self.boton_nueva_categoria = QPushButton("+ Nueva categoría")
        self.boton_nueva_categoria.setToolTip(
            "Crea una categoría y la selecciona para este producto."
        )
        self.boton_nueva_categoria.clicked.connect(
            self.agregar_categoria
        )
        fila_categoria = QHBoxLayout()
        fila_categoria.setContentsMargins(0, 0, 0, 0)
        fila_categoria.addWidget(self.combo_categoria, 1)
        fila_categoria.addWidget(self.boton_nueva_categoria)
        self.combo_proveedor = QComboBox()
        self.campo_costo = QLineEdit()
        self.campo_precio = QLineEdit()
        self.combo_grupo_precio = QComboBox()
        self.boton_nuevo_grupo = QPushButton("+ Nuevo grupo")
        self.boton_nuevo_grupo.clicked.connect(
            self.agregar_grupo_precio
        )
        fila_grupo = QHBoxLayout()
        fila_grupo.addWidget(self.combo_grupo_precio, 1)
        fila_grupo.addWidget(self.boton_nuevo_grupo)
        self.etiqueta_margen = QLabel("—")
        self.etiqueta_margen.setStyleSheet(
            "font-weight: 900; color: #079447;"
        )
        self.campo_costo.textChanged.connect(self._actualizar_margen)
        self.campo_precio.textChanged.connect(self._actualizar_margen)
        self.check_pesable = QCheckBox(
            "Producto pesable · venta por kg"
        )
        self.check_pesable.setObjectName("producto_pesable")
        self.check_pesable.setMinimumHeight(40)
        self.check_pesable.setToolTip(
            "Marcá esta opción cuando la cantidad pueda ingresarse "
            "con decimales o provenga de una balanza."
        )
        self.check_pesable.setStyleSheet(
            """
            QCheckBox#producto_pesable {
                background-color: #eefbf3;
                color: #14532d;
                border: 1px solid #86d5a6;
                border-radius: 7px;
                padding: 8px 12px;
                spacing: 10px;
                font-weight: 800;
            }
            QCheckBox#producto_pesable:hover {
                background-color: #dcf7e7;
                border-color: #12aa52;
            }
            QCheckBox#producto_pesable::indicator {
                width: 20px;
                height: 20px;
                background-color: #ffffff;
                border: 2px solid #64748b;
                border-radius: 4px;
            }
            QCheckBox#producto_pesable::indicator:checked {
                background-color: #12aa52;
                border-color: #0d9245;
            }
            QCheckBox#producto_pesable:disabled {
                background-color: #f1f5f3;
                color: #718079;
                border-color: #cbd5d0;
            }
            """
        )
        self.check_pesable.stateChanged.connect(
            self._actualizar_texto_pesable
        )
        self.check_aplicar_precio_grupo = QCheckBox(
            "Aplicar precio de venta a todo el grupo"
        )
        self.combo_grupo_precio.currentIndexChanged.connect(
            self._actualizar_opcion_precio_grupal
        )

        form.addWidget(QLabel("Código:"), 0, 0)
        form.addWidget(self.campo_codigo, 0, 1)
        form.addWidget(QLabel("PLU:"), 0, 2)
        form.addWidget(self.campo_plu, 0, 3)

        form.addWidget(QLabel("Descripción:"), 1, 0)
        form.addWidget(self.campo_descripcion, 1, 1, 1, 3)

        form.addWidget(QLabel("Categoría:"), 2, 0)
        form.addLayout(fila_categoria, 2, 1)
        form.addWidget(QLabel("Proveedor:"), 2, 2)
        form.addWidget(self.combo_proveedor, 2, 3)

        form.addWidget(QLabel("Costo:"), 3, 0)
        form.addWidget(self.campo_costo, 3, 1)
        form.addWidget(QLabel("Precio de venta:"), 3, 2)
        form.addWidget(self.campo_precio, 3, 3)

        form.addWidget(QLabel("Margen calculado:"), 4, 0)
        form.addWidget(self.etiqueta_margen, 4, 1)
        form.addWidget(QLabel("Grupo de productos:"), 4, 2)
        form.addLayout(fila_grupo, 4, 3)

        form.addWidget(self.check_pesable, 5, 1)
        form.addWidget(
            self.check_aplicar_precio_grupo,
            5,
            3,
        )
        form.setColumnStretch(1, 1)
        form.setColumnStretch(3, 1)
        panel_layout.addLayout(form)

        botones = QHBoxLayout()
        self.boton_guardar = configurar_boton_primario(
            QPushButton("Guardar cambios")
        )
        self.boton_guardar.clicked.connect(self.guardar)
        self.boton_estado = QPushButton("Activar / desactivar")
        self.boton_estado.clicked.connect(self.alternar_estado)
        self.boton_aprobar = QPushButton("Aprobar producto pendiente")
        self.boton_aprobar.clicked.connect(self.aprobar_pendiente)
        botones.addWidget(self.boton_guardar)
        botones.addWidget(self.boton_estado)
        botones.addWidget(self.boton_aprobar)
        botones.addStretch()
        panel_layout.addLayout(botones)

        self.panel_edicion.setVisible(False)
        layout.addWidget(self.panel_edicion)

        if not self.puede_editar:
            layout.addWidget(QLabel("Modo consulta: solo ADMIN y SUPERVISOR pueden modificar productos."))
            self.boton_guardar.setEnabled(False)
            self.boton_estado.setEnabled(False)
            self.boton_importar.setEnabled(False)
            self.boton_nueva_categoria.setEnabled(False)
            self.boton_nuevo_grupo.setEnabled(False)
        self.boton_aprobar.setVisible(self.usuario.rol == "ADMIN")
        self.boton_aprobar.setEnabled(False)

        self._habilitar_formulario(False)

    def _cargar_catalogos(self):
        self.proveedores = listar_proveedores()
        self.grupos_precio = listar_grupos_precio()
        self._recargar_categorias()

        self.combo_proveedor.clear()
        self.combo_proveedor.addItem("Sin proveedor", None)
        for proveedor in self.proveedores:
            self.combo_proveedor.addItem(proveedor.nombre, proveedor.id)

        self.combo_grupo_precio.clear()
        self.combo_grupo_precio.addItem("Sin grupo", None)
        for grupo in self.grupos_precio:
            self.combo_grupo_precio.addItem(grupo.nombre, grupo.id)

    def _recargar_categorias(self, seleccionar_id=None):
        filtro_actual = self.combo_categoria_filtro.currentData()
        categoria_actual = (
            seleccionar_id
            if seleccionar_id is not None
            else self.combo_categoria.currentData()
        )
        self.categorias = listar_categorias()

        self.combo_categoria_filtro.blockSignals(True)
        self.combo_categoria_filtro.clear()
        self.combo_categoria_filtro.addItem("Todas las categorías", None)
        self.combo_categoria.clear()
        for categoria in self.categorias:
            texto = f"{categoria.nombre} · IVA {categoria.iva}"
            self.combo_categoria_filtro.addItem(categoria.nombre, categoria.id)
            self.combo_categoria.addItem(texto, categoria.id)

        if filtro_actual is not None:
            self._seleccionar_por_dato(
                self.combo_categoria_filtro,
                filtro_actual,
            )
        if categoria_actual is not None:
            self._seleccionar_por_dato(
                self.combo_categoria,
                categoria_actual,
            )
        self.combo_categoria_filtro.blockSignals(False)

    def agregar_categoria(self):
        if not self.puede_editar:
            QMessageBox.warning(
                self,
                "Sin permiso",
                "Solo ADMIN o SUPERVISOR pueden crear categorías.",
            )
            return

        dialogo = DialogoNuevaCategoria(self)
        if dialogo.exec() != QDialog.DialogCode.Accepted:
            return
        nombre, iva = dialogo.datos()

        try:
            categoria = crear_categoria(nombre, iva, self.usuario.id)
        except ErrorProducto as error:
            QMessageBox.warning(
                self,
                "No se pudo crear la categoría",
                str(error),
            )
            return

        self._recargar_categorias(categoria.id)
        iva_visible = str(categoria.iva).replace("10.5", "10,5")
        QMessageBox.information(
            self,
            "Categoría creada",
            f"Se creó '{categoria.nombre}' con IVA {iva_visible} y "
            "quedó seleccionada para el producto.",
        )

    def agregar_grupo_precio(self):
        nombre, aceptado = QInputDialog.getText(
            self,
            "Nuevo grupo de productos",
            "Nombre del grupo:",
        )
        if not aceptado:
            return
        try:
            grupo = crear_grupo_precio(nombre, self.usuario.id)
        except ErrorGrupoPrecio as error:
            QMessageBox.warning(
                self,
                "No se pudo crear el grupo",
                str(error),
            )
            return

        self.grupos_precio = listar_grupos_precio()
        self.combo_grupo_precio.clear()
        self.combo_grupo_precio.addItem("Sin grupo", None)
        for grupo_disponible in self.grupos_precio:
            self.combo_grupo_precio.addItem(
                grupo_disponible.nombre,
                grupo_disponible.id,
            )
        self._seleccionar_por_dato(
            self.combo_grupo_precio,
            grupo.id,
        )

    def limpiar_filtros(self):
        self.campo_busqueda.clear()
        self.combo_categoria_filtro.setCurrentIndex(0)
        self.combo_estado.setCurrentText(
            "TODOS" if self.puede_editar else "ACTIVOS"
        )
        self.buscar()

    def exportar_a_excel(self):
        """Exporta a Excel los productos que cumplen los filtros actuales."""
        if not self.puede_exportar:
            QMessageBox.warning(
                self,
                "Sin permiso",
                "Solamente ADMIN o SUPERVISOR pueden exportar productos.",
            )
            return
        if not self.productos:
            QMessageBox.information(
                self,
                "Sin productos",
                "No hay productos para exportar con los filtros actuales.",
            )
            return

        nombre_sugerido = (
            "productos_nepos_"
            + datetime.now().strftime("%Y%m%d_%H%M")
            + ".xlsx"
        )
        ruta, _ = QFileDialog.getSaveFileName(
            self,
            "Exportar catálogo de productos",
            nombre_sugerido,
            "Archivo de Excel (*.xlsx)",
        )
        if not ruta:
            return
        if not ruta.lower().endswith(".xlsx"):
            ruta += ".xlsx"

        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
            from openpyxl.utils import get_column_letter
        except ImportError:
            QMessageBox.critical(
                self,
                "Falta una dependencia",
                "Para exportar instalá openpyxl con:\n"
                "python -m pip install openpyxl",
            )
            return

        try:
            libro = Workbook()
            hoja = libro.active
            hoja.title = "Productos"

            encabezados = [
                "Código",
                "Producto",
                "Categoría",
                "Proveedor",
                "Costo",
                "Precio de venta",
                "Margen (%)",
                "Stock",
                "Grupo",
                "PLU",
                "Activo",
                "Revisión",
            ]
            hoja.append(encabezados)

            for producto in self.productos:
                hoja.append(
                    [
                        producto.codigo,
                        producto.descripcion,
                        (
                            producto.categoria_rel.nombre
                            if producto.categoria_rel
                            else ""
                        ),
                        (
                            producto.proveedor_rel.nombre
                            if producto.proveedor_rel
                            else ""
                        ),
                        float(producto.costo or 0),
                        float(producto.precio or 0),
                        float(producto.incremento or 0) * 100,
                        float(producto.stock or 0),
                        getattr(producto, "grupo_precio_nombre", "") or "",
                        producto.plu or "",
                        "Sí" if producto.activo else "No",
                        (
                            "PENDIENTE"
                            if producto.pendiente_revision
                            else "Aprobado"
                        ),
                    ]
                )

            color_encabezado = "12AA52"
            for celda in hoja[1]:
                celda.fill = PatternFill(
                    fill_type="solid",
                    fgColor=color_encabezado,
                )
                celda.font = Font(color="FFFFFF", bold=True)
                celda.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                )

            for fila in range(2, hoja.max_row + 1):
                hoja.cell(fila, 5).number_format = '$#,##0.00'
                hoja.cell(fila, 6).number_format = '$#,##0.00'
                hoja.cell(fila, 7).number_format = '0.00"%"'
                hoja.cell(fila, 8).number_format = '0.###'

            anchos = [
                18,
                42,
                20,
                24,
                15,
                17,
                14,
                12,
                22,
                12,
                11,
                14,
            ]
            for indice, ancho in enumerate(anchos, start=1):
                hoja.column_dimensions[
                    get_column_letter(indice)
                ].width = ancho

            hoja.freeze_panes = "A2"
            hoja.auto_filter.ref = hoja.dimensions
            hoja.row_dimensions[1].height = 24
            libro.save(ruta)
        except Exception as error:
            QMessageBox.critical(
                self,
                "No se pudo exportar",
                str(error),
            )
            return

        QMessageBox.information(
            self,
            "Exportación terminada",
            f"Se exportaron {len(self.productos)} producto(s) a:\n{ruta}",
        )

    def importar_desde_excel(self):
        if not self.puede_importar:
            QMessageBox.warning(
                self,
                "Sin permiso",
                "Solamente ADMIN o SUPERVISOR pueden importar productos.",
            )
            return
        continuar = QMessageBox.question(
            self,
            "Formato del Excel",
            "COLUMNAS MÍNIMAS RECOMENDADAS\n"
            "A: CODIGO\n"
            "B: DESCRIPCION\n"
            "C: PRECIO DE VENTA\n\n"
            "COLUMNAS OPCIONALES\n"
            "CATEGORIA, IVA, COSTO, STOCK, PROVEEDOR, PESABLE, PLU y GRUPO.\n\n"
            "El orden no es obligatorio: NEPOS reconoce las columnas por "
            "el encabezado. Si faltan datos en un producto nuevo:\n"
            "• Costo y stock quedan en 0.\n"
            "• Categoría queda como SIN CATEGORÍA con IVA 21%.\n"
            "• Proveedor, PLU y grupo quedan vacíos.\n"
            "• PESABLE queda desactivado.\n\n"
            "En una actualización, las columnas opcionales ausentes no "
            "sobrescriben los datos existentes.\n\n"
            "¿Seleccionar el archivo ahora?",
            QMessageBox.StandardButton.Yes
            | QMessageBox.StandardButton.No,
        )
        if continuar != QMessageBox.StandardButton.Yes:
            return
        ruta, _ = QFileDialog.getOpenFileName(
            self,
            "Seleccionar Excel de productos",
            "",
            "Archivos de Excel (*.xlsx *.xlsm)",
        )
        if not ruta:
            return

        try:
            analisis = analizar_excel(ruta)
        except ErrorProducto as error:
            QMessageBox.warning(self, "No se pudo analizar", str(error))
            return

        resumen = (
            f"Filas leídas: {analisis.total_filas}\n"
            f"Productos nuevos: {analisis.nuevos}\n"
            f"Productos existentes: {analisis.existentes}\n"
            f"Categorías a crear: {len(analisis.categorias_nuevas)}\n"
            f"Proveedores a crear: {len(analisis.proveedores_nuevos)}\n"
            f"Duplicados omitidos: {analisis.duplicados_omitidos}\n"
            f"Errores: {len(analisis.errores)}"
        )
        if analisis.errores:
            detalle = "\n".join(analisis.errores[:15])
            if len(analisis.errores) > 15:
                detalle += f"\n... y {len(analisis.errores) - 15} error(es) más."
            QMessageBox.warning(
                self,
                "El Excel contiene errores",
                resumen + "\n\n" + detalle,
            )
            return

        opciones = list(MODOS_IMPORTACION.values())
        opcion, aceptado = QInputDialog.getItem(
            self,
            "Modo de importación",
            resumen + "\n\nElegí qué hacer con los códigos existentes:",
            opciones,
            0,
            False,
        )
        if not aceptado:
            return
        modo = next(clave for clave, nombre in MODOS_IMPORTACION.items() if nombre == opcion)

        confirmacion = QMessageBox.question(
            self,
            "Confirmar importación",
            resumen + f"\n\nModo: {opcion}\n\n¿Continuar?",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
        )
        if confirmacion != QMessageBox.StandardButton.Yes:
            return

        try:
            resultado = importar_excel(ruta, modo, self.usuario.id)
        except ErrorProducto as error:
            QMessageBox.warning(self, "No se pudo importar", str(error))
            return
        except Exception as error:
            QMessageBox.critical(self, "Error de importación", str(error))
            return

        QMessageBox.information(
            self,
            "Importación terminada",
            f"Creados: {resultado['creados']}\n"
            f"Actualizados: {resultado['actualizados']}\n"
            f"Omitidos: {resultado['omitidos']}",
        )
        self.buscar()

    def buscar(self):
        self.productos = listar_productos(
            busqueda=self.campo_busqueda.text(),
            categoria_id=self.combo_categoria_filtro.currentData(),
            estado=self.combo_estado.currentText(),
        )
        self.tabla.setRowCount(len(self.productos))

        for fila, producto in enumerate(self.productos):
            item_codigo = QTableWidgetItem(producto.codigo)
            item_codigo.setData(Qt.ItemDataRole.UserRole, producto.id)
            self.tabla.setItem(fila, 0, item_codigo)
            self.tabla.setItem(fila, 1, QTableWidgetItem(producto.descripcion))
            self.tabla.setItem(fila, 2, QTableWidgetItem(
                producto.categoria_rel.nombre if producto.categoria_rel else ""
            ))
            self.tabla.setItem(fila, 3, QTableWidgetItem(
                producto.proveedor_rel.nombre if producto.proveedor_rel else ""
            ))
            self.tabla.setItem(
                fila, 4, ItemNumerico(
                    formatear_moneda(producto.costo or 0),
                    producto.costo,
                )
            )
            self.tabla.setItem(
                fila, 5, ItemNumerico(
                    formatear_moneda(producto.precio or 0),
                    producto.precio,
                )
            )
            margen = float(producto.incremento or 0) * 100
            self.tabla.setItem(
                fila, 6, ItemNumerico(f"{margen:.2f}%", margen)
            )
            self.tabla.setItem(
                fila, 7, ItemNumerico(
                    formatear_cantidad(producto.stock or 0),
                    producto.stock,
                )
            )
            self.tabla.setItem(
                fila,
                8,
                QTableWidgetItem(
                    getattr(producto, "grupo_precio_nombre", "")
                ),
            )
            self.tabla.setItem(fila, 9, QTableWidgetItem(producto.plu or ""))
            self.tabla.setItem(fila, 10, QTableWidgetItem("Sí" if producto.activo else "No"))
            self.tabla.setItem(
                fila,
                11,
                QTableWidgetItem(
                    "PENDIENTE" if producto.pendiente_revision else "Aprobado"
                ),
            )

        if self._columna_orden is not None:
            self.tabla.sortItems(
                self._columna_orden,
                self._orden_actual,
            )

        self.etiqueta_resultados.setText(f"{len(self.productos)} producto(s)")
        self.producto_id_seleccionado = None
        self.creando_producto = False
        self._limpiar_formulario()
        self._habilitar_formulario(False)
        self.panel_edicion.setVisible(False)

    def _ordenar_por_columna(self, columna):
        if self._columna_orden == columna:
            self._orden_actual = (
                Qt.SortOrder.DescendingOrder
                if self._orden_actual == Qt.SortOrder.AscendingOrder
                else Qt.SortOrder.AscendingOrder
            )
        else:
            self._columna_orden = columna
            self._orden_actual = Qt.SortOrder.AscendingOrder

        self.tabla.sortItems(columna, self._orden_actual)
        self.tabla.horizontalHeader().setSortIndicator(
            columna,
            self._orden_actual,
        )
        self._actualizar_encabezados_orden()

    def _actualizar_encabezados_orden(self):
        for indice, nombre in enumerate(self.COLUMNAS):
            indicador = "↕"
            if indice == self._columna_orden:
                indicador = (
                    "▲"
                    if self._orden_actual == Qt.SortOrder.AscendingOrder
                    else "▼"
                )
            item = self.tabla.horizontalHeaderItem(indice)
            if item is not None:
                item.setText(f"{nombre}  {indicador}")

    def cargar_seleccion(self):
        filas = self.tabla.selectionModel().selectedRows()
        if not filas:
            self.producto_id_seleccionado = None
            self.panel_edicion.setVisible(False)
            self._habilitar_formulario(False)
            return

        fila = filas[0].row()
        producto = self._producto_de_fila(fila)
        if producto is None:
            return
        if not self.puede_editar:
            self.producto_id_seleccionado = producto.id
            self.panel_edicion.setVisible(False)
            return
        self.creando_producto = False
        self.producto_id_seleccionado = producto.id
        self.panel_edicion.setVisible(True)

        self.campo_codigo.setText(producto.codigo)
        self.campo_plu.setText(producto.plu or "")
        self.campo_descripcion.setText(producto.descripcion)
        self._seleccionar_por_dato(self.combo_categoria, producto.categoria_id)
        self._seleccionar_por_dato(self.combo_proveedor, producto.proveedor_id)
        self.campo_costo.setText(str(producto.costo or 0).replace(".", ","))
        self.campo_precio.setText(str(producto.precio or 0).replace(".", ","))
        self._seleccionar_por_dato(
            self.combo_grupo_precio,
            producto.grupo_precio_id,
        )
        self._actualizar_margen()
        self.check_pesable.setChecked(bool(producto.pesable))
        self.check_aplicar_precio_grupo.setChecked(False)
        self.boton_estado.setText("Desactivar" if producto.activo else "Activar")
        self.boton_guardar.setText("Guardar cambios")
        self._habilitar_formulario(self.puede_editar)
        self.boton_aprobar.setEnabled(
            self.usuario.rol == "ADMIN"
            and bool(producto.pendiente_revision)
        )

    @staticmethod
    def _seleccionar_por_dato(combo, valor):
        indice = combo.findData(valor)
        combo.setCurrentIndex(indice if indice >= 0 else 0)

    def _producto_de_fila(self, fila):
        item_codigo = self.tabla.item(fila, 0)
        if item_codigo is None:
            return None
        producto_id = item_codigo.data(Qt.ItemDataRole.UserRole)
        return next(
            (
                producto
                for producto in self.productos
                if producto.id == producto_id
            ),
            None,
        )

    def _leer_numero(self, campo, nombre):
        try:
            return convertir_decimal_finito(
                campo.text(),
                nombre=nombre,
                minimo=0,
                maximo=100000000,
                interpretar_punto_miles=True,
            )
        except ValueError as error:
            raise ErrorProducto(str(error)) from error

    def _actualizar_margen(self):
        try:
            costo = self._leer_numero(self.campo_costo, "El costo")
            precio = self._leer_numero(
                self.campo_precio, "El precio"
            )
            margen = ((precio / costo) - 1) * 100 if costo > 0 else 0
            self.etiqueta_margen.setText(f"{margen:.2f}%")
            color = "#079447" if margen >= 0 else "#b42318"
            self.etiqueta_margen.setStyleSheet(
                f"font-weight:900; color:{color};"
            )
        except ErrorProducto:
            self.etiqueta_margen.setText("—")

    def guardar(self):
        if not self.creando_producto and not self.producto_id_seleccionado:
            QMessageBox.information(self, "Elegí un producto", "Seleccioná una fila primero.")
            return

        if self.creando_producto:
            try:
                crear_producto(
                    usuario_id=self.usuario.id,
                    codigo=self.campo_codigo.text(),
                    plu=self.campo_plu.text(),
                    descripcion=self.campo_descripcion.text(),
                    categoria_id=self.combo_categoria.currentData(),
                    proveedor_id=self.combo_proveedor.currentData(),
                    costo=self._leer_numero(self.campo_costo, "El costo"),
                    precio=self._leer_numero(
                        self.campo_precio, "El precio de venta"
                    ),
                    grupo_precio_id=self.combo_grupo_precio.currentData(),
                    pesable=self.check_pesable.isChecked(),
                )
            except ErrorProducto as error:
                QMessageBox.warning(self, "No se pudo crear", str(error))
                return

            QMessageBox.information(
                self,
                "Producto creado",
                "El producto se creó activo, con stock inicial 0.",
            )
            self.buscar()
            return

        aplicar_precio_grupo = (
            self.check_aplicar_precio_grupo.isChecked()
        )
        if aplicar_precio_grupo:
            grupo = self.combo_grupo_precio.currentText()
            respuesta = QMessageBox.question(
                self,
                "Actualizar precio del grupo",
                f"Se aplicará el nuevo precio de venta a todos los "
                f"productos activos del grupo '{grupo}'.\n\n"
                "El costo de cada producto no se modificará y su margen "
                "se recalculará automáticamente.\n\n¿Continuar?",
                QMessageBox.StandardButton.Yes
                | QMessageBox.StandardButton.No,
            )
            if respuesta != QMessageBox.StandardButton.Yes:
                return

        try:
            resultado = actualizar_producto(
                producto_id=self.producto_id_seleccionado,
                usuario_id=self.usuario.id,
                codigo=self.campo_codigo.text(),
                plu=self.campo_plu.text(),
                descripcion=self.campo_descripcion.text(),
                categoria_id=self.combo_categoria.currentData(),
                proveedor_id=self.combo_proveedor.currentData(),
                costo=self._leer_numero(self.campo_costo, "El costo"),
                precio=self._leer_numero(
                    self.campo_precio, "El precio de venta"
                ),
                grupo_precio_id=self.combo_grupo_precio.currentData(),
                pesable=self.check_pesable.isChecked(),
                aplicar_precio_grupo=aplicar_precio_grupo,
            )
        except ErrorProducto as error:
            QMessageBox.warning(self, "No se pudo guardar", str(error))
            return

        cantidad_grupo = (
            resultado.get("productos_grupo_actualizados", 0)
            if resultado
            else 0
        )
        mensaje = "Los cambios se guardaron correctamente."
        if cantidad_grupo:
            mensaje += (
                f"\n\nPrecio actualizado en {cantidad_grupo} "
                "producto(s) del grupo."
            )
        QMessageBox.information(
            self,
            "Producto actualizado",
            mensaje,
        )
        self.buscar()

    def alternar_estado(self):
        filas = self.tabla.selectionModel().selectedRows()
        if not filas:
            QMessageBox.information(self, "Elegí un producto", "Seleccioná una fila primero.")
            return

        producto = self._producto_de_fila(filas[0].row())
        if producto is None:
            return
        nuevo_estado = not bool(producto.activo)
        accion = "activar" if nuevo_estado else "desactivar"
        respuesta = QMessageBox.question(
            self,
            "Confirmar cambio",
            f"¿Querés {accion} el producto '{producto.descripcion}'?",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
        )
        if respuesta != QMessageBox.StandardButton.Yes:
            return

        try:
            cambiar_estado_producto(producto.id, nuevo_estado, self.usuario.id)
        except ErrorProducto as error:
            QMessageBox.warning(self, "No se pudo actualizar", str(error))
            return

        self.buscar()

    def aprobar_pendiente(self):
        if not self.producto_id_seleccionado:
            QMessageBox.information(
                self,
                "Elegí un producto",
                "Seleccioná una fila primero.",
            )
            return
        respuesta = QMessageBox.question(
            self,
            "Aprobar producto",
            "Antes de aprobar, verificá categoría, proveedor, costo y precio. "
            "¿Confirmás la aprobación?",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
        )
        if respuesta != QMessageBox.StandardButton.Yes:
            return
        try:
            aprobar_producto(
                self.producto_id_seleccionado,
                self.usuario.id,
            )
        except ErrorProductoPendiente as error:
            QMessageBox.warning(
                self,
                "No se pudo aprobar",
                str(error),
            )
            return
        QMessageBox.information(
            self,
            "Producto aprobado",
            "El producto dejó de estar pendiente de revisión.",
        )
        self.buscar()

    def _limpiar_formulario(self):
        self.campo_codigo.clear()
        self.campo_plu.clear()
        self.campo_descripcion.clear()
        self.campo_costo.clear()
        self.campo_precio.clear()
        self.combo_categoria.setCurrentIndex(0)
        self.combo_proveedor.setCurrentIndex(0)
        self.combo_grupo_precio.setCurrentIndex(0)
        self.etiqueta_margen.setText("—")
        self.check_pesable.setChecked(False)
        self.check_aplicar_precio_grupo.setChecked(False)

    def _actualizar_opcion_precio_grupal(self, _indice=None):
        disponible = (
            self.puede_editar
            and self.producto_id_seleccionado is not None
            and not self.creando_producto
            and self.combo_grupo_precio.currentData() is not None
        )
        self.check_aplicar_precio_grupo.setEnabled(disponible)
        if not disponible:
            self.check_aplicar_precio_grupo.setChecked(False)

    def _actualizar_texto_pesable(self, _estado=None):
        prefijo = "✓ " if self.check_pesable.isChecked() else ""
        self.check_pesable.setText(
            f"{prefijo}Producto pesable · venta por kg"
        )

    def _habilitar_formulario(self, habilitado):
        for control in (
            self.campo_codigo, self.campo_plu, self.campo_descripcion,
            self.combo_categoria, self.combo_proveedor, self.campo_costo,
            self.campo_precio, self.combo_grupo_precio,
            self.check_pesable,
        ):
            control.setEnabled(habilitado)
        self.check_aplicar_precio_grupo.setEnabled(
            habilitado
            and self.puede_editar
            and self.combo_grupo_precio.currentData() is not None
        )
        if not self.check_aplicar_precio_grupo.isEnabled():
            self.check_aplicar_precio_grupo.setChecked(False)
        self.boton_guardar.setEnabled(habilitado and self.puede_editar)
        self.boton_nueva_categoria.setEnabled(
            habilitado and self.puede_editar
        )
        self.boton_nuevo_grupo.setEnabled(
            habilitado and self.puede_editar
        )
        self.boton_estado.setEnabled(
            self.producto_id_seleccionado is not None and self.puede_editar
        )
        if not habilitado:
            self.boton_aprobar.setEnabled(False)

    def nuevo_producto(self):
        if not self.puede_editar:
            return
        self.tabla.clearSelection()
        self.producto_id_seleccionado = None
        self.creando_producto = True
        self._limpiar_formulario()
        self.panel_edicion.setVisible(True)
        self._habilitar_formulario(True)
        self.check_aplicar_precio_grupo.setEnabled(False)
        self.boton_estado.setEnabled(False)
        self.boton_aprobar.setEnabled(False)
        self.boton_guardar.setText("Crear producto")
        self.campo_codigo.setFocus()